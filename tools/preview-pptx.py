#!/usr/bin/env python3
"""معاينة مرئية للـ PPTX: يقرأ الأشكال الفعلية من الملف ويرسمها HTML للتصوير.

ليست بديلاً عن فتح الملف في باوربوينت، لكنها تكشف الفيض والتداخل وسوء
المحاذاة لأنها تقرأ إحداثيات الملف نفسه — لا تعيد اشتقاقها من كود التوليد.
"""
import base64
import html
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "out/instagram/pptx-preview.html"
SLIDE_W, SLIDE_H = 1080, 1350


def px(v):
    return Emu(v).inches * 96


NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _srgb(el):
    """أول لون srgbClr داخل عنصر XML، وإلا None."""
    if el is None:
        return None
    node = el.find(f".//{{{NS_A}}}srgbClr")
    return "#" + node.get("val") if node is not None else None


def _sp_pr(sh):
    """spPr الخاص بالشكل — بحث مقيّد حتى لا نلتقط ألوان النصوص."""
    for tag in ("spPr", "grpSpPr"):
        for ns in (
            "http://schemas.openxmlformats.org/presentationml/2006/main",
            "http://schemas.openxmlformats.org/drawingml/2006/main",
        ):
            el = sh._element.find(f"{{{ns}}}{tag}")
            if el is not None:
                return el
    return None


def shape_fill(sh):
    pr = _sp_pr(sh)
    if pr is None:
        return None
    fill = pr.find(f"{{{NS_A}}}solidFill")  # ابن مباشر فقط، لا داخل a:ln
    return _srgb(fill)


def shape_line(sh):
    pr = _sp_pr(sh)
    if pr is None:
        return None
    ln = pr.find(f"{{{NS_A}}}ln")
    if ln is None or ln.find(f"{{{NS_A}}}noFill") is not None:
        return None
    return _srgb(ln)


def shape_geom(sh):
    pr = _sp_pr(sh)
    if pr is None:
        return None
    g = pr.find(f"{{{NS_A}}}prstGeom")
    return g.get("prst") if g is not None else None


NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def slide_bg(slide):
    for path in (f"{{{NS_P}}}cSld/{{{NS_P}}}bg", f"{{{NS_P}}}bg"):
        el = slide._element.find(path)
        if el is not None:
            return _srgb(el)
    return None


def data_uri(image):
    return "data:%s;base64,%s" % (
        image.content_type,
        base64.b64encode(image.blob).decode(),
    )


def render_shape(sh):
    x, y = px(sh.left), px(sh.top)
    w, h = px(sh.width), px(sh.height)
    base = f"left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"

    if sh.shape_type == 13:  # PICTURE
        return (
            f'<img class="s" style="{base}object-fit:contain" '
            f'src="{data_uri(sh.image)}">'
        )

    style = base
    geom = shape_geom(sh)
    fill = shape_fill(sh)
    line = shape_line(sh)
    if fill:
        style += f"background:{fill};"
    if line:
        style += f"border:1.4px solid {line};"
    if geom == "ellipse":
        style += "border-radius:50%;"
    elif geom == "diamond":
        style += "transform:rotate(45deg);"

    inner = ""
    if sh.has_text_frame and sh.text_frame.text.strip():
        rows = []
        for para in sh.text_frame.paragraphs:
            spans = []
            for r in para.runs:
                f = r.font
                size = f.size.pt * (96 / 72) if f.size else 24
                col = _srgb(r._r.find(f"{{{NS_A}}}rPr")) or "#2E2622"
                fam = f.name or "Tajawal"
                s = (
                    f"font-family:'{fam}';font-size:{size:.1f}px;color:{col};"
                    f"font-weight:{'700' if f.bold else '400'};"
                )
                if f.underline:
                    s += "text-decoration:underline;"
                spans.append(
                    f'<span style="{s}">{html.escape(r.text)}</span>'
                )
            rows.append("<div>%s</div>" % ("".join(spans) or "&nbsp;"))
        inner = (
            '<div class="tx" style="direction:rtl">%s</div>' % "".join(rows)
        )

    return f'<div class="s" style="{style}">{inner}</div>'


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else str(
        ROOT / "out/instagram/lomond-instagram-editable.pptx"
    )
    prs = Presentation(path)
    fonts = ROOT / "public/fonts"
    faces = "".join(
        f"@font-face{{font-family:'{fam}';src:url('file://{fonts}/{f}');"
        f"font-weight:{wt};}}"
        for fam, f, wt in [
            ("Amiri", "Amiri-Bold.ttf", 700),
            ("Amiri", "Amiri-Regular.ttf", 400),
            ("Tajawal", "Tajawal-Bold.ttf", 700),
            ("Tajawal", "Tajawal-Medium.ttf", 400),
            ("Times New Roman", "Tinos-Bold.ttf", 700),
            ("Times New Roman", "Tinos-Regular.ttf", 400),
        ]
    )

    slides = []
    for i, slide in enumerate(prs.slides, 1):
        bg = slide_bg(slide) or "#EEE4DA"
        body = "".join(render_shape(sh) for sh in slide.shapes)
        slides.append(
            f'<div class="slide" style="background:{bg}">{body}'
            f'<div class="num">{i}</div></div>'
        )

    OUT.write_text(
        "<style>%s body{margin:0;background:#888;display:grid;"
        "grid-template-columns:repeat(3,1fr);gap:8px}"
        ".slide{position:relative;width:%dpx;height:%dpx;overflow:hidden}"
        ".s{position:absolute;box-sizing:border-box}"
        ".tx{width:100%%;height:100%%;display:flex;flex-direction:column;"
        "align-items:center;justify-content:center;text-align:center;"
        "line-height:1.35}"
        ".num{position:absolute;left:6px;top:6px;font:700 26px monospace;"
        "color:#0008}</style>%s"
        % (faces, SLIDE_W, SLIDE_H, "".join(slides)),
        encoding="utf-8",
    )
    print(OUT)


if __name__ == "__main__":
    main()
